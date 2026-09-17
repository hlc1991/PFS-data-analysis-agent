import json
import shutil
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from agent.agent import BusinessAgent
from data.sources.csv import CSVDataSource
from infrastructure.paths import data_path

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "data" / "fixtures" / "pfs_sales.csv"


class PfsExportArtifactTests(unittest.TestCase):
    def setUp(self):
        self.tmp = data_path("outputs", "exports", "_pfs-export-test")
        shutil.rmtree(self.tmp, ignore_errors=True)
        self.tmp.mkdir(parents=True, exist_ok=True)
        self.agent = BusinessAgent(
            client=None, model="pfs-local-test",
            data_source=CSVDataSource(str(FIXTURE), "pfs_sales.csv"),
            session_id="pfs-export-test",
        )
        self.agent._get_export_dir = lambda: str(self.tmp)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_excel_is_parseable(self):
        result = self.agent._tool_export_excel(["pfs_sales"], "pfs_sales_export")
        self.assertIn("/api/export/", result)
        from openpyxl import load_workbook
        path = self.tmp / "pfs_sales_export.xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["pfs_sales"])
        rows = list(workbook["pfs_sales"].iter_rows(values_only=True))
        self.assertEqual(rows[0], ("month", "region", "product", "sales_amount"))
        self.assertEqual(len(rows), 10)
        workbook.close()

    def test_word_report_is_parseable(self):
        result = self.agent._tool_export_report(
            "PFS 销售分析报告",
            [{"heading": "核心结论", "content": "华东销售额最高。"}],
        )
        self.assertIn("/api/export/", result)
        from docx import Document
        files = list(self.tmp.glob("*.docx"))
        self.assertEqual(len(files), 1)
        text = "\n".join(p.text for p in Document(files[0]).paragraphs)
        self.assertIn("PFS 销售分析报告", text)
        self.assertIn("核心结论", text)
        self.assertIn("华东销售额最高", text)
        docx_path = next(self.tmp.glob("*.docx"))
        with zipfile.ZipFile(docx_path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
            styles_xml = archive.read("word/styles.xml").decode("utf-8")
        self.assertIn('w:eastAsia="Hiragino Sans"', document_xml)
        self.assertIn('w:eastAsia="Hiragino Sans"', styles_xml)

    def test_ppt_is_parseable(self):
        result = self.agent._tool_generate_ppt(
            "PFS 销售分析",
            [{"layout": "cover", "params": {"title": "PFS 销售分析"}},
             {"layout": "closing", "params": {"title": "结论", "message": "完成"}}],
            "pfs_sales",
        )
        if result.startswith("❌ PPT 模块加载失败"):
            self.skipTest(result)
        self.assertIn("/api/export/", result)
        from pptx import Presentation
        files = list(self.tmp.glob("*.pptx"))
        self.assertEqual(len(files), 1)
        presentation = Presentation(files[0])
        self.assertEqual(len(presentation.slides), 2)
        texts = " ".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("PFS 销售分析", texts)
        with zipfile.ZipFile(files[0]) as archive:
            slide_xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        self.assertIn('typeface="Hiragino Sans"', slide_xml)

    def test_dashboard_html_export_is_pfs_branded(self):
        import api.dashboard as dashboard_module
        from api.dashboard import build_dashboard
        with patch.object(dashboard_module, "_DASHBOARD_DIR", str(self.tmp)):
            data = build_dashboard(
                self.agent.data_source, self.agent._chart_store,
                session_id="pfs-export-test", workspace_id="", name="PFS 销售看板",
                widgets_spec=[{
                    "id": "sales-by-region", "title": "地区销售额", "chart_type": "Bar_Chart",
                    "sql": "SELECT region, SUM(sales_amount) AS total_sales FROM pfs_sales GROUP BY region",
                    "field_mapping": {"x": "region", "y": "total_sales"},
                }], color_scheme="pfs",
            )
            dashboard = json.loads(
                (self.tmp / (data["dashboard_id"] + ".json")).read_text(encoding="utf-8")
            )
        from api.dashboard_html_export import build_export_html
        html = build_export_html(dashboard, self.agent._chart_store)
        self.assertIn("PFS 数据分析 Agent", html)
        self.assertNotIn("智析 Agent", html)
        self.assertIn("PFS 销售看板", html)


if __name__ == "__main__":
    unittest.main()
